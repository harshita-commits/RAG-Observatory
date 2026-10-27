import importlib.util
import io
import unittest
from observatory.documents import extract, chunks_from_files


def text_pdf():
    content = b"BT /F1 12 Tf 50 700 Td (Annual leave is 24 days.) Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n".encode()+content+b"\nendstream"]
    data, offsets = b"%PDF-1.4\n", [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += f"{index} 0 obj\n".encode()+obj+b"\nendobj\n"
    xref = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n"
    data += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    return data+f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()


class FormatTests(unittest.TestCase):
    def test_csv_json_html_and_unsupported_inputs(self):
        self.assertIn("days: 24", extract("a.csv", b"name,days\nleave,24")[0]["text"])
        self.assertEqual(extract("a.jsonl", b'{"x":1}\n{"x":2}')[1]["item"], 2)
        html = extract("a.html", b"<script>secret()</script><style>hidden</style><p>Policy: 24 days</p>")[0]["text"]
        self.assertNotIn("secret", html)
        self.assertIn("24 days", html)
        for name, content in (("a.exe", b"x"), ("a.txt", b"")):
            with self.assertRaises(ValueError): extract(name, content)

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "optional PDF parser missing")
    def test_real_pdf_text_and_explicit_ocr_requirement(self):
        from pypdf import PdfWriter
        self.assertIn("24 days", extract("a.pdf", text_pdf())[0]["text"])
        self.assertEqual(extract("a.pdf", text_pdf())[0]["page"], 1)
        writer = PdfWriter()
        writer.add_blank_page(100,100)
        data = io.BytesIO()
        writer.write(data)
        with self.assertRaisesRegex(ValueError, "OCR"):
            extract("scan.pdf", data.getvalue())

    @unittest.skipUnless(importlib.util.find_spec("docx"), "optional DOCX parser missing")
    def test_real_docx_paragraphs_and_tables(self):
        from docx import Document
        document = Document()
        document.add_paragraph("Leave entitlement")
        table = document.add_table(rows=1, cols=2)
        table.cell(0,0).text, table.cell(0,1).text = "Annual", "24 days"
        data = io.BytesIO()
        document.save(data)
        result = extract("a.docx", data.getvalue())
        self.assertIn("Leave", result[0]["text"])
        self.assertIn("24 days", result[1]["text"])

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "optional XLSX parser missing")
    def test_real_xlsx_preserves_sheet_and_row(self):
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.title = "Policies"
        workbook.active.append(["Policy", "Days"])
        workbook.active.append(["Annual leave", 24])
        data = io.BytesIO()
        workbook.save(data)
        workbook.close()
        result = extract("a.xlsx", data.getvalue())[0]
        self.assertEqual(result["sheet"], "Policies")
        self.assertEqual(result["row"], 2)
        self.assertIn("Days: 24", result["text"])

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "optional PPTX parser missing")
    def test_real_pptx_preserves_slide_number(self):
        from pptx import Presentation
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = "Annual leave is 24 days"
        data = io.BytesIO()
        deck.save(data)
        result = extract("a.pptx", data.getvalue())[0]
        self.assertEqual(result["slide"], 1)
        self.assertIn("24 days", result["text"])

    def test_duplicate_upload_names_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "Duplicate"):
            chunks_from_files([("a.txt", b"one"), ("a.txt", b"two")], "dataset")


if __name__ == "__main__": unittest.main()
