"""Local, text-oriented parsers. No fetching URLs, executing macros, or OCR."""
import csv
import hashlib
import io
import json
import zipfile
from html.parser import HTMLParser
from pathlib import Path

FORMATS = {".txt", ".md", ".csv", ".tsv", ".json", ".jsonl", ".html", ".htm", ".pdf", ".docx", ".pptx", ".xlsx"}
MAX_FILE_BYTES = 10 * 1024 * 1024


class TextHTMLParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts, self.hidden = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}: self.hidden += 1

    def handle_endtag(self, tag):
        if tag in {"script", "style"}: self.hidden = max(self.hidden-1, 0)

    def handle_data(self, data):
        if not self.hidden: self.parts.append(data)


def extract(name, content):
    suffix = Path(name).suffix.lower()
    if suffix not in FORMATS:
        raise ValueError(f"Unsupported format {suffix or '(none)'}. Supported: {', '.join(sorted(FORMATS))}")
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("Each file must be at most 10 MiB")
    if suffix in {".docx", ".xlsx", ".pptx"}:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 10000 or sum(i.file_size for i in archive.infolist()) > 64*1024*1024:
                raise ValueError("Office archive exceeds the local decompression limit")
    sections = []
    def add(text, **location):
        if text and text.strip(): sections.append({"text": text.strip(), **location})
    try:
        if suffix in {".txt", ".md"}:
            add(content.decode("utf-8-sig"))
        elif suffix in {".html", ".htm"}:
            parser = TextHTMLParser()
            parser.feed(content.decode("utf-8-sig"))
            add(" ".join(parser.parts))
        elif suffix in {".csv", ".tsv"}:
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")), delimiter="\t" if suffix == ".tsv" else ",")
            for row_number, row in enumerate(reader, 2):
                add("; ".join(f"{key}: {value}" for key, value in row.items()), row=row_number)
        elif suffix in {".json", ".jsonl"}:
            text = content.decode("utf-8-sig")
            rows = [json.loads(line) for line in text.splitlines() if line.strip()] if suffix == ".jsonl" else [json.loads(text)]
            for number, value in enumerate(rows, 1):
                add(json.dumps(value, ensure_ascii=False), item=number)
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("Encrypted PDFs must be decrypted before import")
            empty = []
            for page_number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if not text.strip(): empty.append(page_number)
                add(text, page=page_number)
            # No silent partial ingestion of scanned/mixed evidence.
            if empty:
                raise ValueError(f"PDF has pages without extractable text: {empty[:10]}. Apply OCR first; no OCR is bundled.")
        elif suffix == ".docx":
            from docx import Document
            from docx.table import Table
            document = Document(io.BytesIO(content))
            for block in document.iter_inner_content():
                add("\n".join(" | ".join(cell.text for cell in row.cells) for row in block.rows) if isinstance(block, Table) else block.text)
        elif suffix == ".pptx":
            from pptx import Presentation
            deck = Presentation(io.BytesIO(content))
            for number, slide in enumerate(deck.slides, 1):
                text = []
                for shape in slide.shapes:
                    if shape.has_text_frame: text.append(shape.text_frame.text)
                    if shape.has_table:
                        text.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
                add("\n".join(text), slide=number)
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in workbook:
                    rows = sheet.iter_rows(values_only=True)
                    headers = next(rows, ())
                    for number, values in enumerate(rows, 2):
                        if any(v is not None for v in values):
                            add("; ".join(f"{headers[i] if i < len(headers) and headers[i] is not None else 'column '+str(i+1)}: {v}" for i, v in enumerate(values) if v is not None), sheet=sheet.title, row=number)
            finally:
                workbook.close()
    except ImportError as exc:
        raise ValueError("This format needs optional parsers. Install: python -m pip install -r requirements-documents.txt") from exc
    if not sections:
        raise ValueError(f"No extractable text in {name}. Images, scanned content and empty files need preprocessing.")
    return sections


def chunks_from_files(files, name, chunk_words=300, overlap_words=50):
    if not isinstance(chunk_words, int) or not isinstance(overlap_words, int) or chunk_words < 1 or not 0 <= overlap_words < chunk_words:
        raise ValueError("Require chunk_words > overlap_words >= 0")
    if not files or len(files) > 100:
        raise ValueError("Import between 1 and 100 files per batch")
    if sum(len(content) for _, content in files) > 20*1024*1024:
        raise ValueError("Total file content must be at most 20 MiB")
    documents, names = [], set()
    for source, content in files:
        source = source.replace("\\", "/")
        if source in names: raise ValueError(f"Duplicate filename: {source}")
        names.add(source)
        file_hash = hashlib.sha256(source.encode()).hexdigest()[:12]
        for section_index, section in enumerate(extract(source, content)):
            words = section["text"].split()
            for start in range(0, len(words), chunk_words-overlap_words):
                end = min(start+chunk_words, len(words))
                identifier = f"doc-{file_hash}-s{section_index}-{start:06d}"
                documents.append({"id": identifier, "topic": "custom", "source": source,
                                  "word_start": start, "word_end": end, "section": section_index,
                                  "location": {k: v for k, v in section.items() if k != "text"}, "text": " ".join(words[start:end])})
                if len(documents) > 10000: raise ValueError("Local benchmark limit: 10,000 chunks")
                if end == len(words): break
    return {"name": name, "ingestion": {"chunk_words": chunk_words, "overlap_words": overlap_words,
            "formats": sorted({Path(f).suffix.lower() for f, _ in files})}, "documents": documents, "cases": []}
